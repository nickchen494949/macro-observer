import openpyxl
wb = openpyxl.load_workbook('scratch/Yardeni_Sector_Operating_EPS.xlsm', data_only=True)
ws = wb.worksheets[0]
for r in ws.iter_rows(min_row=8, max_row=14, values_only=True): print(r)
print('...')
for r in ws.iter_rows(min_row=ws.max_row-6, max_row=ws.max_row, values_only=True): print(r)
ws = wb.worksheets[1]
for r in ws.iter_rows(min_row=9, max_row=14, values_only=True): print(r)
import collections
print(collections.Counter(r[4] for r in ws.iter_rows(min_row=9, values_only=True)))
print(collections.Counter(r[8] for r in ws.iter_rows(min_row=9, values_only=True)))
print(sorted(set((r[1], r[2]) for r in ws.iter_rows(min_row=9, values_only=True))))
