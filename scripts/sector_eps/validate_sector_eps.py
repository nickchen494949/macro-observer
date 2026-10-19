#!/usr/bin/env python3
"""Validation for sector EPS extraction. Usage: python3 validate_sector_eps.py <xlsx> <out_dir>

Only the internal-consistency check is automated. (An attempted multpl.com cross-check was discarded:
multpl publishes inflation-adjusted EPS, which is not comparable with S&P nominal EPS.)
"""
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

xlsx, out = Path(sys.argv[1]), Path(sys.argv[2])
q = pd.read_csv(out / "sector_eps_quarterly.csv")
wb = openpyxl.load_workbook(xlsx, data_only=True)
ws = wb["SECTOR EPS"]

# ---- 1. internal consistency: sum of 4 actual quarters == the sheet's own annual 'YYYY EPS' column
def annual_block(title_row_text):
    tr = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value).strip() == title_row_text)
    h = next(r for r in range(tr - 2, tr + 3) if ws.cell(r, 1).value == "INDEX NAME")
    cols = {}
    for c in range(3, ws.max_column + 1):
        m = re.match(r"^(\d{4})\s+EPS$", str(ws.cell(h, c).value or "").strip())
        if m:
            cols[int(m[1])] = c
    res = {}
    r = tr + 1
    for r in range(tr + 1, tr + 14):
        name = str(ws.cell(r, 1).value or "").strip()
        if name.startswith("S&P 500"):
            sec = re.sub(r"\(proforma.*\)", "", name[len("S&P 500"):]).strip() or "ALL (index)"
            res[sec] = {y: ws.cell(r, c).value for y, c in cols.items()}
    return res

rows = []
for basis, title in (("operating", "Operating Earnings Per Share by Economic Sector"),
                     ("as_reported", "As Reported Earnings Per Share by Economic Sector")):
    ann = annual_block(title)
    d = q[(q["index"] == "S&P 500") & (q.basis == basis) & (q.status == "actual")].copy()
    d["year"] = d.period.str[:4].astype(int)
    for (sec, yr), g in d.groupby(["sector", "year"]):
        if len(g) == 4 and isinstance(ann.get(sec, {}).get(yr), (int, float)):
            rows.append(dict(basis=basis, sector=sec, year=yr, sum_q=g.eps.sum(), annual=ann[sec][yr],
                             diff=g.eps.sum() - ann[sec][yr]))
v = pd.DataFrame(rows)
v.to_csv(out / "validation_annual_vs_quarters.csv", index=False)
print("== 1. sum of 4 quarters vs sheet annual EPS ==")
print(v.groupby("basis")["diff"].agg(n="size", max_abs=lambda s: s.abs().max(), mean_abs=lambda s: s.abs().mean()))
bad = v[v["diff"].abs() > 0.05]
print("cells with |diff|>0.05:", len(bad)); print(bad.head(15).to_string(index=False))
