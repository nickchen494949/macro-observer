#!/usr/bin/env python3
"""
Parse S&P Dow Jones Indices' official `sp-500-eps-est.xlsx` (sheet 'SECTOR EPS') into tidy CSV.

Source (free, official): https://www.spglobal.com/spdji/en/documents/additional-material/sp-500-eps-est.xlsx
  - The live URL blocks scripts (HTTP 403); download it in a browser, or use a Wayback capture.
Usage:
  python3 build_sector_eps.py <path-to-xlsx> [out_dir]

Output (long format):
  sector_eps_quarterly.csv : index, sector, basis(operating|as_reported), period (YYYYQn),
                             period_end, eps, status(actual|estimate)
  sector_eps_wide_<basis>.csv : one column per sector (actuals only)
Columns S&P labels "YYYYE Qn" are ESTIMATES and are flagged status=estimate (never mixed into the
wide actual tables). NOTE: operating EPS is finalised later than as-reported EPS, so the last actual
quarter differs by basis.
"""
import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd

QRE = re.compile(r"^(\d{4})(E?)\s*Q([1-4])$")
SECTORS = {
    "Consumer Discretionary": "Consumer Discretionary",
    "Consumer Staples": "Consumer Staples",
    "Energy": "Energy",
    "Financials": "Financials",
    "Health Care": "Health Care",
    "Industrials": "Industrials",
    "Information Technology": "Information Technology",
    "Materials": "Materials",
    "Communication Services": "Communication Services",
    "Utilities": "Utilities",
    "Real Estate": "Real Estate",
}
INDEXES = ["S&P 500", "S&P 400", "S&P 600", "S&P Composite 1500"]


def parse(xlsx: Path) -> tuple[pd.DataFrame, pd.Timestamp]:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    ws = wb["SECTOR EPS"]
    actuals_through = pd.Timestamp(ws.cell(3, 2).value)
    as_of = pd.Timestamp(ws.cell(2, 2).value)

    # locate the two blocks by their titles in column A
    block_rows, basis = {}, None
    titles = {"Operating Earnings Per Share by Economic Sector": "operating",
              "As Reported Earnings Per Share by Economic Sector": "as_reported"}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and v.strip() in titles:
            basis = titles[v.strip()]
            block_rows[basis] = {"title_row": r, "rows": []}
        elif basis and isinstance(v, str) and v.strip() not in ("INDEX NAME",):
            block_rows[basis]["rows"].append(r)

    rows = []
    for basis, info in block_rows.items():
        # header row = the 'INDEX NAME' row adjacent to the title row (before it in block 1, after in block 2)
        h = next(r for r in range(info["title_row"] - 2, info["title_row"] + 3)
                 if ws.cell(r, 1).value == "INDEX NAME")
        qcols = {}
        for c in range(3, ws.max_column + 1):
            m = QRE.match(str(ws.cell(h, c).value or "").strip())
            if m and c not in qcols.values() and (int(m[1]), int(m[3])) not in qcols:
                # S&P labels estimate columns with an "E" (e.g. "2025E Q3"); keep that flag explicitly
                qcols[(int(m[1]), int(m[3]))] = (c, m[2] == "E")
        for r in info["rows"]:
            name = str(ws.cell(r, 1).value).strip()
            idx = next((i for i in sorted(INDEXES, key=len, reverse=True) if name.startswith(i)), None)
            if idx is None:
                continue
            sector = name[len(idx):].strip()
            sector = re.sub(r"\(proforma.*\)", "", sector).strip() or "ALL (index)"
            for (y, q), (c, is_est) in qcols.items():
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)):
                    pe = pd.Period(f"{y}Q{q}").end_time.normalize()
                    rows.append(dict(index=idx, sector=sector, basis=basis, period=f"{y}Q{q}",
                                     period_end=pe.date(), eps=float(v),
                                     status="estimate" if is_est else "actual"))
    df = pd.DataFrame(rows)
    df.attrs["as_of"], df.attrs["actuals_through"] = as_of, actuals_through
    return df, actuals_through


def main():
    xlsx = Path(sys.argv[1])
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else xlsx.parent
    out.mkdir(parents=True, exist_ok=True)
    df, thr = parse(xlsx)
    df.to_csv(out / "sector_eps_quarterly.csv", index=False)
    for basis in ("operating", "as_reported"):
        w = df[(df["index"] == "S&P 500") & (df.basis == basis) & (df.status == "actual")]
        wide = w.pivot(index="period", columns="sector", values="eps").sort_index()
        wide.to_csv(out / f"sector_eps_wide_{basis}.csv")
    import json
    last = {b: df[(df.basis == b) & (df.status == "actual")].period.max() for b in ("operating", "as_reported")}
    first_est = {b: (df[(df.basis == b) & (df.status == "estimate")].period.min()
                     if ((df.basis == b) & (df.status == "estimate")).any() else None)
                 for b in ("operating", "as_reported")}
    meta = dict(source="S&P Dow Jones Indices, sp-500-eps-est.xlsx (sheet SECTOR EPS)",
                source_url="https://www.spglobal.com/spdji/en/documents/additional-material/sp-500-eps-est.xlsx",
                file_as_of=str(df.attrs["as_of"].date()),
                header_actuals_through=str(thr.date()),
                last_actual_quarter=last, first_estimate_quarter=first_est,
                note="Sector EPS is per sector-index unit; levels are NOT additive across sectors.")
    (out / "sector_eps_meta.json").write_text(json.dumps(meta, indent=2))
    print(f"file as of {meta['file_as_of']} | last actual quarter: {last} | rows={len(df)}")


if __name__ == "__main__":
    main()
